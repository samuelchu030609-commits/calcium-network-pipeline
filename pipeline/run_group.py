"""Aggregate several per-recording *_metrics.xlsx into one group-comparison workbook.

Each recording's workbook already contains a machine-readable **"Summary"** sheet
(long format: Section | Metric | Value). This tool reads that sheet from every
recording, lines the headline metrics up side by side, and computes per-group
descriptive stats and (for a two-group design) an exploratory contrast — so you
can compare conditions (e.g. WT vs KCNT1) without re-deriving anything in Excel.

It reads only the finished workbooks and writes a new one; it never touches the
per-recording analysis, so it is safe to run any time.

Grouping is the USER's call — genotype is not encoded in the files. Each recording
is assigned a group label you provide (the GUI lets you edit it per row; the CLI
guesses from the folder name and you can override). Nothing here infers WT vs mutant.

Contract (kept stable for the GUI):

    build_comparison(items, out_path) -> str
      items    : list of dicts {"path": <xlsx>, "group": <label>, "label": <name>}
      out_path : where to write group_comparison.xlsx
      returns  : the written path

CLI:
    python -m pipeline.run_group --scan PARENT [--out OUT.xlsx]
    python -m pipeline.run_group WB1.xlsx:GroupA WB2.xlsx:GroupA WB3.xlsx:GroupB
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path

import openpyxl

# Curated headline metrics pulled from each recording's "Summary" sheet, in order.
# (section, metric) must match the Summary sheet's first two columns exactly.
# kind: "int" | "num" | "pct" | "text" — controls formatting and whether it is a
# numeric metric that gets group stats + a contrast.
SUMMARY_METRICS = [
    ("Recording", "n_cells_analyzed",                 "Cells analyzed",              "int"),
    ("Recording", "n_active_cells",                    "Active cells",                "int"),
    ("Recording", "percent_active_cells",             "% active cells",              "pct"),
    ("Recording", "duration_min",                      "Duration (min)",              "num"),
    ("Recording", "fps",                               "Frame rate (Hz)",             "num"),
    ("Events",    "event_signal_source",               "Event source",                "text"),
    ("Events",    "total_events_detected",             "Total events",                "int"),
    ("Event rate (active,/min)", "mean",               "Event rate (active, /min)",   "num"),
    ("Synchrony - STTC (primary)", "STTC_z_vs_chance", "STTC z vs chance (primary)",  "num"),
    ("Synchrony - STTC (primary)", "STTC_excess_over_chance", "STTC excess over chance", "num"),
    ("Synchrony - STTC (primary)", "mean_STTC",        "STTC mean (raw)",             "num"),
    ("Synchrony - Pearson (secondary)", "mean_pearson_dff", "Pearson mean (secondary)", "num"),
    ("Population coupling (Okun)", "mean_population_coupling", "Population coupling",   "num"),
    ("Functional connectivity graph (FDR)", "edge_density", "FC edge density",         "num"),
    ("Functional connectivity graph (FDR)", "mean_degree",  "FC mean degree",          "num"),
    ("Network bursts", "n_bursts_shuffle_p99",         "Network bursts (sig.)",       "int"),
    ("Network bursts", "max_coactive_fraction",        "Max coactive fraction",       "num"),
    ("Network bursts", "fraction_cells_in_significant_bursts", "Cells in sig. bursts", "num"),
    ("Cell assemblies (PCA+MP+ICA)", "n_cells_in_an_assembly", "Cells in assemblies",  "int"),
]

# The numeric metrics that get group means/SD and a contrast (event source excluded).
NUMERIC_KINDS = {"int", "num", "pct"}


class GroupError(Exception):
    """User-facing problem building the comparison."""


def _to_number(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_summary(xlsx_path) -> dict:
    """Read a recording's 'Summary' sheet into {(section, metric): raw_value}."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.is_file():
        raise GroupError(f"No such workbook: {xlsx_path}")
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:
        raise GroupError(f"Could not open {xlsx_path.name}: {exc}")
    if "Summary" not in wb.sheetnames:
        raise GroupError(
            f"{xlsx_path.name} has no 'Summary' sheet — is it a metrics workbook "
            f"from this pipeline?")
    out = {}
    ws = wb["Summary"]
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 3:
            continue
        section, metric, value = row[0], row[1], row[2]
        if section in (None, "Section") or metric is None:
            continue
        out[(str(section), str(metric))] = value
    wb.close()
    return out


def _recording_name(xlsx_path) -> str:
    """A short recording id: strip a trailing _metrics(.*) from the file stem."""
    stem = Path(xlsx_path).stem
    for suffix in ("_metrics",):
        i = stem.find(suffix)
        if i != -1:
            return stem[:i] or stem
    return stem


def _mean_sd_median(values):
    xs = [v for v in values if v is not None]
    n = len(xs)
    if n == 0:
        return {"n": 0, "mean": None, "sd": None, "median": None}
    mean = sum(xs) / n
    if n >= 2:
        var = sum((x - mean) ** 2 for x in xs) / (n - 1)  # sample SD
        sd = var ** 0.5
    else:
        sd = None
    srt = sorted(xs)
    mid = n // 2
    median = srt[mid] if n % 2 else (srt[mid - 1] + srt[mid]) / 2
    return {"n": n, "mean": mean, "sd": sd, "median": median}


def _mann_whitney_p(a, b):
    """Two-sided Mann-Whitney U p-value, or None if unavailable/underpowered."""
    a = [x for x in a if x is not None]
    b = [x for x in b if x is not None]
    if len(a) < 3 or len(b) < 3:
        return None
    try:
        from scipy.stats import mannwhitneyu
        return float(mannwhitneyu(a, b, alternative="two-sided").pvalue)
    except Exception:
        return None


def detection_of(xlsx_path) -> tuple[str, tuple | None]:
    """(readable label, comparable signature) of the cell-detection settings a
    workbook's recording was made with, or ("unknown", None).

    Looks beside the workbook: in plane0 (detection_settings.json, else ops.npy), or,
    for a copy in RESULTS/, at RESULTS/<id>_detection_settings.json and then the
    recording's own plane0.
    """
    import json
    p = Path(xlsx_path)
    rid = p.name[:-len("_metrics.xlsx")] if p.name.endswith("_metrics.xlsx") else p.stem
    if p.parent.name == "RESULTS":
        jsons = [p.parent / f"{rid}_detection_settings.json",
                 p.parent.parent / rid / "suite2p" / "plane0" / "detection_settings.json"]
        plane0 = p.parent.parent / rid / "suite2p" / "plane0"
    else:
        jsons = [p.parent / "detection_settings.json"]
        plane0 = p.parent
    for j in jsons:
        try:
            d = json.loads(j.read_text(encoding="utf-8"))
            sig = (round(float(d["diameter_px"]), 3), float(d["cellprob_threshold"]),
                   float(d["flow_threshold"]), d["img"])
            return str(d.get("profile") or "unnamed settings"), sig
        except Exception:
            continue
    try:
        import numpy as np
        ops = np.load(plane0 / "ops.npy", allow_pickle=True).item()
        cs = ops["detection"]["cellpose_settings"]
        sig = (round(float(np.atleast_1d(ops["diameter"])[-1]), 3),
               float(cs["cellprob_threshold"]), float(cs["flow_threshold"]), cs["img"])
        return (f"{sig[0]:g} px, cellprob {sig[1]:g}, flow {sig[2]:g}, {sig[3]}", sig)
    except Exception:
        return "unknown", None


def build_comparison(items: list, out_path) -> str:
    """Aggregate `items` (each {path, group, label}) into a comparison workbook."""
    if not items:
        raise GroupError("No recordings to compare.")

    # --- read every recording's Summary ---
    rows = []          # one dict per recording: label, group, + display->value
    event_sources = set()
    detections = {}    # signature -> label
    for it in items:
        summ = read_summary(it["path"])
        det_label, det_sig = detection_of(it["path"])
        if det_sig is not None:
            detections.setdefault(det_sig, det_label)
        rec = {"label": it.get("label") or _recording_name(it["path"]),
               "group": str(it.get("group") or "ungrouped"),
               "detection": det_label}
        for section, metric, display, kind in SUMMARY_METRICS:
            raw = summ.get((section, metric))
            if kind in NUMERIC_KINDS:
                rec[display] = _to_number(raw)
            else:
                rec[display] = None if raw is None else str(raw)
        src = rec.get("Event source")
        if src:
            event_sources.add(src)
        rows.append(rec)

    groups = []
    for r in rows:
        if r["group"] not in groups:
            groups.append(r["group"])

    # methodological guard: event-based metrics are not comparable across event
    # sources (CASCADE calibrated spikes vs raw dF/F0). Flag it rather than hide it.
    mixed_sources = len(event_sources) > 1
    # Same guard for cell detection: different settings find different cells.
    mixed_detection = len(detections) > 1

    wb = openpyxl.Workbook()

    # ---- Sheet 1: Per recording (wide) ----
    ws = wb.active
    ws.title = "Per recording"
    header = ["Recording", "Group"] + [d for _, _, d, _ in SUMMARY_METRICS] + ["Cell detection"]
    ws.append(header)
    for r in rows:
        ws.append([r["label"], r["group"]] +
                  [r.get(d) for _, _, d, _ in SUMMARY_METRICS] + [r["detection"]])

    # ---- Sheet 2: By group (descriptive stats) ----
    wsg = wb.create_sheet("By group")
    wsg.append(["Metric", "Group", "n", "mean", "SD", "median"])
    for section, metric, display, kind in SUMMARY_METRICS:
        if kind not in NUMERIC_KINDS:
            continue
        for g in groups:
            vals = [r[display] for r in rows if r["group"] == g]
            s = _mean_sd_median(vals)
            wsg.append([display, g, s["n"], s["mean"], s["sd"], s["median"]])

    # ---- Sheet 3: Contrast (only meaningful for exactly two groups) ----
    wsc = wb.create_sheet("Contrast")
    if len(groups) == 2:
        gA, gB = groups
        wsc.append([f"Two-group contrast: {gB} minus {gA}"])
        wsc.append(["Metric", f"{gA} mean (n)", f"{gB} mean (n)",
                    f"difference ({gB}-{gA})", "Mann-Whitney p (exploratory)"])
        for section, metric, display, kind in SUMMARY_METRICS:
            if kind not in NUMERIC_KINDS:
                continue
            va = [r[display] for r in rows if r["group"] == gA]
            vb = [r[display] for r in rows if r["group"] == gB]
            sa, sb = _mean_sd_median(va), _mean_sd_median(vb)
            diff = (sb["mean"] - sa["mean"]) if (sa["mean"] is not None and sb["mean"] is not None) else None
            p = _mann_whitney_p(va, vb)
            wsc.append([
                display,
                f"{sa['mean']:.4g} (n={sa['n']})" if sa["mean"] is not None else f"— (n={sa['n']})",
                f"{sb['mean']:.4g} (n={sb['n']})" if sb["mean"] is not None else f"— (n={sb['n']})",
                round(diff, 6) if diff is not None else None,
                round(p, 4) if p is not None else "n too small (need >=3/group)",
            ])
    else:
        wsc.append([f"Contrast needs exactly 2 groups; you have {len(groups)} "
                    f"({', '.join(groups)}). See 'By group' for per-group stats."])

    # ---- Sheet 4: How to read ----
    wsh = wb.create_sheet("How to read")
    notes = [
        ["What this workbook is"],
        ["One row per recording (Per recording), per-group means +/- SD (By group),"],
        ["and, for a two-group design, an exploratory contrast (Contrast)."],
        [""],
        ["Grouping"],
        ["Groups are the labels you assigned; genotype is not read from the files."],
        [""],
        ["Reading the numbers"],
        ["Primary synchrony is 'STTC z vs chance' (not raw mean STTC) — report that."],
        ["SD is the sample standard deviation; 'n' is the number of recordings."],
        ["The Mann-Whitney p is two-sided, unpaired, and EXPLORATORY; it is only"],
        ["computed with >=3 recordings per group and does not correct for testing"],
        ["many metrics. Treat small-n comparisons as descriptive."],
    ]
    if mixed_sources:
        notes += [
            [""],
            ["!! WARNING: event sources are mixed across these recordings:"],
            ["   " + ", ".join(sorted(event_sources))],
            ["   Event-based metrics (event rate, STTC, bursts) are NOT directly"],
            ["   comparable between CASCADE (calibrated spikes) and dF/F0 (dye)."],
            ["   Compare like-with-like event sources."],
        ]
    if mixed_detection:
        notes += [
            [""],
            ["!! WARNING: these recordings were detected with DIFFERENT cell-detection settings:"],
            *[["   " + lbl] for lbl in sorted(detections.values())],
            ["   Different settings find different cells, so cell counts, % active and"],
            ["   synchrony are NOT directly comparable. Analyse every recording of an"],
            ["   experiment with the same detection profile."],
        ]
    for line in notes:
        wsh.append(line)

    out_path = str(out_path)
    wb.save(out_path)
    return out_path


# ─────────────────────────── CLI ───────────────────────────
def _guess_group(path) -> str:
    """Default group label from the folder path (lab naming). User can override."""
    s = str(path).lower()
    if "fluo" in s:
        return "Fluo-4"
    if any(k in s for k in ("srs10", "ss10", "8f")):
        return "jGCaMP8f"
    if any(k in s for k in ("srs9", "ss9", "8s")):
        return "jGCaMP8s"
    return "ungrouped"


def _scan_workbooks(parent) -> list:
    """Every recording's workbook under `parent`, each recording ONCE.

    tools/analyze_folder.py keeps the original in <rec>/suite2p/plane0/ and a copy in
    RESULTS/; counting both would double n. Set-aside earlier runs
    (suite2p_previous_*) are not current results and are skipped.
    """
    parent = Path(parent)
    found = {}
    for root, dirs, files in os.walk(parent):
        dirs[:] = [d for d in dirs
                   if not d.startswith(".") and not d.startswith("suite2p_previous")]
        here = Path(root)
        for f in files:
            if f.startswith("~$"):            # Excel's lock file while a workbook is open
                continue
            if f.endswith("_metrics.xlsx") or (f.endswith(".xlsx") and "_metrics" in f):
                if here.name == "RESULTS":
                    key, prefer = (here.parent, f), False
                elif here.name == "plane0" and here.parent.name == "suite2p":
                    key, prefer = (here.parent.parent.parent, f), True
                else:
                    key, prefer = (here, f), True
                if key not in found or prefer:
                    found[key] = here / f
    return sorted(found.values())


def main() -> None:
    ap = argparse.ArgumentParser(description="Aggregate *_metrics.xlsx into a group comparison.")
    ap.add_argument("workbooks", nargs="*",
                    help="workbook paths, optionally WB.xlsx:GroupLabel")
    ap.add_argument("--scan", metavar="PARENT",
                    help="scan a folder for *_metrics.xlsx (group guessed from path)")
    ap.add_argument("--out", default=None, help="output path (default: ./group_comparison.xlsx)")
    args = ap.parse_args()

    items = []
    for spec in args.workbooks:
        if ":" in spec and not spec[1:3] == ":\\":  # allow Windows drive letters
            path, group = spec.rsplit(":", 1)
        else:
            path, group = spec, None
        items.append({"path": path, "group": group or _guess_group(path)})
    if args.scan:
        for wb in _scan_workbooks(args.scan):
            items.append({"path": str(wb), "group": _guess_group(wb)})

    if not items:
        ap.error("give some workbooks, or --scan PARENT")

    out = args.out or "group_comparison.xlsx"
    path = build_comparison(items, out)
    print(f"[run_group] wrote {path} ({len(items)} recordings, "
          f"groups: {', '.join(sorted({i['group'] for i in items}))})")


if __name__ == "__main__":
    main()
