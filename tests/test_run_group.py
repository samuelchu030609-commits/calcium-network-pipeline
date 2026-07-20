"""Portable tests for pipeline.run_group (no real data, no network).

Run directly:   python tests/test_run_group.py
Or with pytest: pytest tests/test_run_group.py
"""
import os
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline.run_group import build_comparison, read_summary, GroupError, _guess_group


def _make_metrics_workbook(path, *, pct_active, sttc_z, event_source="CASCADE discrete spikes (calibrated)"):
    """Write a minimal workbook with just the 'Summary' sheet the tool reads."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Section", "Metric", "Value"])
    ws.append(["Recording", "n_cells_analyzed", 100])
    ws.append(["Recording", "n_active_cells", int(pct_active)])
    ws.append(["Recording", "percent_active_cells", pct_active])
    ws.append(["Recording", "fps", 45])
    ws.append(["Recording", "duration_min", 1.0])
    ws.append(["Events", "event_signal_source", event_source])
    ws.append(["Events", "total_events_detected", 50])
    ws.append(["Synchrony - STTC (primary)", "STTC_z_vs_chance", sttc_z])
    wb.save(path)


def test_build_comparison_two_groups():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        a1, a2, a3 = d / "a1.xlsx", d / "a2.xlsx", d / "a3.xlsx"
        b1, b2, b3 = d / "b1.xlsx", d / "b2.xlsx", d / "b3.xlsx"
        _make_metrics_workbook(a1, pct_active=10, sttc_z=2.0)
        _make_metrics_workbook(a2, pct_active=20, sttc_z=4.0)
        _make_metrics_workbook(a3, pct_active=30, sttc_z=6.0)  # A mean %active = 20
        _make_metrics_workbook(b1, pct_active=40, sttc_z=8.0)
        _make_metrics_workbook(b2, pct_active=50, sttc_z=10.0)
        _make_metrics_workbook(b3, pct_active=60, sttc_z=12.0)  # B mean %active = 50

        items = [{"path": str(p), "group": "WT"} for p in (a1, a2, a3)] + \
                [{"path": str(p), "group": "KCNT1"} for p in (b1, b2, b3)]
        out = d / "cmp.xlsx"
        build_comparison(items, out)
        assert out.exists()

        wb = openpyxl.load_workbook(out, data_only=True)
        assert wb.sheetnames[:4] == ["Per recording", "By group", "Contrast", "How to read"]

        # Per recording: 6 data rows + 1 header
        assert wb["Per recording"].max_row == 7

        # By group: check WT/KCNT1 means for % active cells (20 and 50)
        by_group = {(r[0], r[1]): r for r in wb["By group"].iter_rows(values_only=True)}
        wt = by_group[("% active cells", "WT")]
        kc = by_group[("% active cells", "KCNT1")]
        assert wt[2] == 3 and abs(wt[3] - 20.0) < 1e-9, f"WT %active mean wrong: {wt}"
        assert kc[2] == 3 and abs(kc[3] - 50.0) < 1e-9, f"KCNT1 %active mean wrong: {kc}"

        # Contrast: difference (KCNT1 - WT) for % active = 30; a p-value is computed (n=3 each)
        contrast = list(wb["Contrast"].iter_rows(values_only=True))
        assert "Two-group contrast" in str(contrast[0][0])
        row = next(r for r in contrast if r[0] == "% active cells")
        assert abs(row[3] - 30.0) < 1e-9, f"contrast diff wrong: {row}"
        assert isinstance(row[4], (int, float)), f"expected a numeric p (n>=3/group): {row[4]}"
    print("PASS test_build_comparison_two_groups")


def test_mixed_event_source_warns():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        casc = d / "casc.xlsx"; dye = d / "dye.xlsx"
        _make_metrics_workbook(casc, pct_active=10, sttc_z=2.0,
                               event_source="CASCADE discrete spikes (calibrated)")
        _make_metrics_workbook(dye, pct_active=80, sttc_z=100.0, event_source="dF/F0")
        out = d / "cmp.xlsx"
        build_comparison([{"path": str(casc), "group": "g1"},
                          {"path": str(dye), "group": "g2"}], out)
        wb = openpyxl.load_workbook(out)
        how = "\n".join(str(r[0]) for r in wb["How to read"].iter_rows(values_only=True) if r and r[0])
        assert "WARNING" in how and "event sources are mixed" in how, how
    print("PASS test_mixed_event_source_warns")


def test_missing_summary_sheet_errors():
    with tempfile.TemporaryDirectory() as d:
        bad = Path(d) / "bad.xlsx"
        wb = openpyxl.Workbook(); wb.active.title = "NotSummary"; wb.save(bad)
        try:
            read_summary(bad)
        except GroupError:
            print("PASS test_missing_summary_sheet_errors")
            return
        raise AssertionError("expected GroupError for a workbook with no Summary sheet")


def test_guess_group():
    assert _guess_group("x/SS9_KOLFs/2169_metrics.xlsx") == "jGCaMP8s"
    assert _guess_group("x/SS10_thing/6_metrics.xlsx") == "jGCaMP8f"
    assert _guess_group("x/Fluo4AM_y/1_metrics.xlsx") == "Fluo-4"
    assert _guess_group("x/random/rec.xlsx") == "ungrouped"
    print("PASS test_guess_group")


if __name__ == "__main__":
    test_build_comparison_two_groups()
    test_mixed_event_source_warns()
    test_missing_summary_sheet_errors()
    test_guess_group()
    print("\nALL run_group TESTS PASSED")
